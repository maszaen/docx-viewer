import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# ==========================================
# HALAMAN 1: JUDUL & PENDAHULUAN
# ==========================================
title = doc.add_heading('Laporan Preview Dokumen Cicool CI3', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

p_info = doc.add_paragraph()
p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_author = p_info.add_run('Halaman 1 dari 3 | Modul Client-Side Viewer | Status: Read-Only')
run_author.font.italic = True
run_author.font.color.rgb = RGBColor(120, 120, 120)

doc.add_paragraph()

doc.add_heading('1. Ringkasan Eksekutif', level=1)
doc.add_paragraph(
    'Modul ini dirancang untuk menampilkan pratinjau dokumen Microsoft Word (.docx) secara langsung '
    'di dalam peramban web (browser) tanpa menggunakan layanan pihak ketiga seperti Google Docs Viewer '
    'maupun Office Online Viewer. Hal ini menjaga kerahasiaan dan privasi dokumen internal sistem Cicool.'
)

doc.add_heading('2. Arsitektur Client-Side', level=1)
doc.add_paragraph(
    'Dengan memindahkan tugas rendering dari server PHP ke peramban pengguna, beban CPU dan memori server '
    'dapat dihemat secara drastis. Berkas OpenXML dibaca dan diurai secara lokal menggunakan pustaka JavaScript murni.'
)

# --- PAGE BREAK KE HALAMAN 2 ---
doc.add_page_break()

# ==========================================
# HALAMAN 2: KOMPONEN & TABEL SPESIFIKASI
# ==========================================
h2 = doc.add_heading('3. Komponen dan Library Pendukung', level=1)
doc.add_paragraph('Implementasi dilakukan dengan JavaScript murni (Vanilla JS) dengan dependensi berikut:')

bullets = [
    ('JSZip (v3.10.1):', ' Bertanggung jawab untuk membaca dan mengekstrak struktur ZIP berkas .docx.'),
    ('docx-preview (v0.3.3):', ' Merender XML dokumen OpenXML menjadi elemen HTML & styling CSS.'),
    ('Query String & IndexedDB:', ' Mekanisme transmisi data antara halaman utama dan tab pratinjau.')
]

for label, text in bullets:
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(label)
    r.bold = True
    p.add_run(text)

doc.add_heading('4. Tabel Spesifikasi Kompatibilitas', level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'

hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Format File'
hdr_cells[1].text = 'Status Dukungan'
hdr_cells[2].text = 'Keterangan Teknis'

for cell in hdr_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

data = [
    ('.docx (Word 2007+)', 'Didukung Penuh', 'Format berbasis OpenXML ZIP terkompresi'),
    ('.doc (Word 97-2003)', 'Ditolak Otomatis', 'Format biner legacy CFBF tidak didukung'),
    ('.pdf / .txt', 'Format Terpisah', 'Memerlukan modul viewer terpisah')
]

for ext, status, note in data:
    row_cells = table.add_row().cells
    row_cells[0].text = ext
    row_cells[1].text = status
    row_cells[2].text = note

# --- PAGE BREAK KE HALAMAN 3 ---
doc.add_page_break()

# ==========================================
# HALAMAN 3: BATASAN & CATATAN TEKNIS
# ==========================================
doc.add_heading('5. Batasan Teknis (Technical Limitations)', level=1)
p_limit = doc.add_paragraph()
r_warn = p_limit.add_run('PERHATIAN: ')
r_warn.bold = True
r_warn.font.color.rgb = RGBColor(180, 50, 50)
p_limit.add_run(
    'Rendering di sisi klien (client-side) menggunakan docx-preview bersifat read-only. '
    'Tata letak halaman mungkin tidak 100% pixel-perfect dibandingkan aplikasi desktop Microsoft Word asli, '
    'terutama terkait pagination otomatis, font khusus sistem operasi, header/footer dinamis, dan makro Office.'
)

doc.add_heading('6. Kesimpulan dan Rekomendasi', level=1)
doc.add_paragraph(
    'Pendekatan client-side ini sangat direkomendasikan untuk sistem CodeIgniter 3 / Cicool '
    'karena kestabilan, portabilitas tanpa build tools, serta kepatuhan penuh terhadap privasi data dokumen.'
)

output_path = 'h:/VSCode/docx-preview-cicool/sample.docx'
doc.save(output_path)
print(f'Sample docx berhasil dibuat ulang dengan 3 halaman terpisah: {output_path}')
